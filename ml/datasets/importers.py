"""Lossless source readers. Register future document/OCR readers through IMPORTERS."""

import csv
import io
import json
import zipfile
from pathlib import Path
from typing import Protocol


class Importer(Protocol):
    def __call__(self, data: bytes) -> list[dict]: ...


def delimited(data, delimiter=None):
    text = data.decode("utf-8-sig")
    if not text.strip():
        raise ValueError("The file is empty")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|").delimiter
        except csv.Error:
            raise ValueError("TXT requires delimited columns with a header, e.g. headword<TAB>definition")
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)):
        raise ValueError("Column headers must be unique")
    return [{str(k) if k is not None else "__extra_cells__": v for k, v in row.items()} for row in reader]


def json_rows(data):
    def invalid_constant(value):
        raise ValueError(f"Non-finite JSON number is not supported: {value}")

    value = json.loads(data, parse_constant=invalid_constant)
    if isinstance(value, dict):
        value = value.get("entries", value)
        if isinstance(value, dict):
            value = [{"headword": k, "definition": v} for k, v in value.items()]
    if not isinstance(value, list) or any(not isinstance(x, dict) for x in value):
        raise ValueError(
            "JSON must contain an array of objects, an entries array, or a word-to-meaning object"
        )
    return value


def xlsx_rows(data):
    from openpyxl import load_workbook

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Expanded spreadsheet exceeds 100 MB")
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    try:
        rows = book.active.iter_rows(values_only=True)
        headers = [str(v) if v is not None else "" for v in next(rows)]
        if not all(headers) or len(set(headers)) != len(headers):
            raise ValueError("Spreadsheet column headers must be non-empty and unique")
        result = []
        for values in rows:
            if not any(v is not None for v in values):
                continue
            result.append(
                dict(zip(headers, [v.isoformat() if hasattr(v, "isoformat") else v for v in values]))
            )
            if len(result) > 100000:
                raise ValueError("Spreadsheet exceeds 100,000 rows")
        return result
    finally:
        book.close()


IMPORTERS: dict[str, Importer] = {
    ".csv": lambda b: delimited(b, ","),
    ".tsv": lambda b: delimited(b, "\t"),
    ".txt": delimited,
    ".json": json_rows,
    ".xlsx": xlsx_rows,
}
ALIASES = {
    "headword": ["headword", "word", "indigenous_word", "term", "lemma"],
    "definitions": ["definitions", "definition", "english", "meaning", "translation", "gloss"],
    "part_of_speech": ["part_of_speech", "pos", "word_class"],
    "examples": ["examples", "example", "sentence"],
    "dialect": ["dialect"],
    "alternate_spellings": ["alternate_spellings", "variants", "variant", "spellings"],
    "notes": ["notes", "note"],
    "source": ["source", "reference"],
}


def parse_file(filename, data, max_rows=100000):
    ext = Path(filename).suffix.lower()
    if ext not in IMPORTERS:
        raise ValueError("Supported formats: CSV, TSV, XLSX, JSON and delimited TXT")
    try:
        rows = IMPORTERS[ext](data)
    except (UnicodeError, json.JSONDecodeError, csv.Error, zipfile.BadZipFile, StopIteration) as exc:
        raise ValueError(f"Cannot read {ext} file: {exc}") from exc
    if not rows:
        raise ValueError("No dictionary rows were found")
    if len(rows) > max_rows:
        raise ValueError(f"Maximum {max_rows:,} rows per upload")
    columns = list(dict.fromkeys(k for row in rows for k in row if k != "__extra_cells__"))
    lookup = {k.strip().lower().replace(" ", "_"): k for k in columns}
    mapping = {
        field: next((lookup[a] for a in aliases if a in lookup), "") for field, aliases in ALIASES.items()
    }
    return rows, columns, {k: v for k, v in mapping.items() if v}
